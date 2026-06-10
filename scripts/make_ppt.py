# -*- coding: utf-8 -*-
"""生成两页项目介绍 PPT（精简版，突出实用性）。运行：python scripts/make_ppt.py"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RUNS = os.path.join(ROOT, "data", "runs", "tool_yolo")
OUT = os.path.join(ROOT, "智能工具检测系统_项目介绍.pptx")

BG     = RGBColor(0x1A, 0x1A, 0x2E)
CARD   = RGBColor(0x24, 0x24, 0x40)
ACCENT = RGBColor(0x4C, 0xC9, 0xF0)
GREEN  = RGBColor(0x52, 0xD2, 0x73)
ORANGE = RGBColor(0xFF, 0xB3, 0x4D)
WHITE  = RGBColor(0xF2, 0xF2, 0xF7)
GREY   = RGBColor(0x9A, 0x9A, 0xB2)
FONT = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]


def bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def rect(slide, l, t, w, h, color, radius=True):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def txt(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=6, line_spacing=1.0):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(4)
    tf.margin_top = tf.margin_bottom = Pt(2)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (sstr, sz, c, b) in para:
            r = p.add_run()
            r.text = sstr
            r.font.size = Pt(sz)
            r.font.color.rgb = c
            r.font.bold = b
            r.font.name = FONT
    return tb


# ================= Slide 1 — 它能做什么（价值） =================
s = prs.slides.add_slide(blank)
bg(s)
rect(s, 0, 0, SW, Inches(0.18), ACCENT, radius=False)

txt(s, Inches(0.7), Inches(0.6), Inches(8.0), Inches(1.6), [
    [("智能工具检测系统", 42, WHITE, True)],
    [("摄像头扫一眼，自动清点电工工具、即时提醒缺件", 19, ACCENT, False)],
], space_after=10)

# 三个实用价值 chip（极简）
chips = [
    ("🎥 实时清点", "同屏多件工具一次框全", ACCENT),
    ("🔔 缺件提醒", "少了哪件，语音报出来", GREEN),
    ("💻 即装即用", "纯 CPU，打包成 EXE 双击运行", ORANGE),
]
cw = Inches(3.85)
cgap = Inches(0.25)
ctop = Inches(2.5)
for i, (h, b, c) in enumerate(chips):
    left = Inches(0.7) + i * (cw + cgap)
    rect(s, left, ctop, cw, Inches(1.25), CARD)
    rect(s, left, ctop, cw, Inches(0.1), c, radius=False)
    txt(s, left, ctop + Inches(0.22), cw, Inches(1.0), [
        [(h, 19, WHITE, True)],
        [(b, 13, GREY, False)],
    ], align=PP_ALIGN.CENTER, space_after=5)

# 一句应用场景
txt(s, Inches(0.7), Inches(4.05), Inches(12), Inches(0.5),
    [[("适用：电工作业前后工具清点、工位监控、实训考核——防止工具遗漏、错拿。",
       14, GREY, False)]])

# hero 检测图
hero = os.path.join(RUNS, "val_batch0_pred.jpg")
txt(s, Inches(0.7), Inches(4.7), Inches(12), Inches(0.4),
    [[("▍真实检测效果", 13, ACCENT, True)]])
s.shapes.add_picture(hero, Inches(0.7), Inches(5.1), height=Inches(2.0))

prs.save(OUT)  # 占位保存，后面继续追加


# ================= Slide 2 — 为什么好用（效果 + 底气） =================
s = prs.slides.add_slide(blank)
bg(s)
rect(s, 0, 0, SW, Inches(0.18), ACCENT, radius=False)

txt(s, Inches(0.7), Inches(0.55), Inches(12), Inches(0.9),
    [[("识别准、跑得动、好部署", 34, WHITE, True)]])

# 指标卡（4 个，大数字）
metrics = [("mAP@50", "0.91", GREEN), ("精确率", "0.88", ACCENT),
           ("召回率", "0.89", ACCENT), ("推理节奏", "0.4s", ORANGE)]
mw = Inches(2.7)
mgap = Inches(0.25)
mtop = Inches(1.7)
for i, (k, v, c) in enumerate(metrics):
    left = Inches(0.7) + i * (mw + mgap)
    rect(s, left, mtop, mw, Inches(1.5), CARD)
    txt(s, left, mtop + Inches(0.18), mw, Inches(0.85),
        [[(v, 38, c, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, left, mtop + Inches(1.0), mw, Inches(0.4),
        [[(k, 14, GREY, False)]], align=PP_ALIGN.CENTER)

# 三条"好用"理由（精简）
reasons = [
    ("不挑工具", "光滑无纹理的工具也能认——深度学习看整体外观，突破了特征匹配的上限", GREEN),
    ("不挑机器", "YOLO11s 纯 CPU 训练与运行，普通电脑即可，无需显卡", ACCENT),
    ("不挑人用", "模型缺失自动回退、双击即开；缺件实时语音播报，看一眼就知道少了啥", ORANGE),
]
rtop = Inches(3.55)
for i, (h, b, c) in enumerate(reasons):
    top = rtop + i * Inches(0.92)
    rect(s, Inches(0.7), top, Inches(7.3), Inches(0.8), CARD)
    rect(s, Inches(0.7), top, Inches(0.1), Inches(0.8), c, radius=False)
    txt(s, Inches(0.95), top + Inches(0.1), Inches(6.9), Inches(0.65), [
        [(h, 15, c, True)],
        [(b, 11.5, GREY, False)],
    ], space_after=1, line_spacing=1.0)

# 右侧训练曲线图
res = os.path.join(RUNS, "results.png")
txt(s, Inches(8.3), Inches(3.2), Inches(4.6), Inches(0.4),
    [[("▍训练收敛稳定", 13, ACCENT, True)]])
s.shapes.add_picture(res, Inches(8.3), Inches(3.6), width=Inches(4.6))

txt(s, Inches(0.7), Inches(7.0), Inches(12), Inches(0.3),
    [[("YOLO11s · 9 类电工工具 · 验证集真实指标（data/runs/tool_yolo）", 10.5, GREY, False)]])

prs.save(OUT)
print("已生成：", OUT)
